"""Word (.docx) and Excel (.xlsx) claim artifacts."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.shared import Pt, RGBColor
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill


def _doc(carrier: str, heading: str) -> Document:
    d = Document()
    h = d.add_paragraph()
    r = h.add_run(carrier)
    r.bold = True
    r.font.size = Pt(15)
    r.font.color.rgb = RGBColor(0x14, 0x28, 0x50)
    sub = d.add_paragraph().add_run(heading)
    sub.font.size = Pt(11)
    sub.font.color.rgb = RGBColor(0x5A, 0x5A, 0x5A)
    d.add_paragraph()
    return d


def coverage_letter_docx(cid: str, c: dict, out: Path) -> Path:
    d = _doc(c["carrier"], "RESERVATION OF RIGHTS / COVERAGE POSITION LETTER")
    d.add_paragraph(f"Re: Claim {cid} — Insured {c['insured']['name']} — Date of Loss {c['loss']['date']}")
    d.add_paragraph(
        f"Dear {c['insured']['name']},\n\n{c['carrier']} is investigating the above-referenced claim under a full "
        "reservation of rights. Coverage is afforded under the Personal Auto Policy, subject to all terms, conditions, "
        "exclusions, and limitations. Liability limits applicable to this loss are $250,000 per person / $500,000 per "
        "occurrence for bodily injury and $100,000 for property damage.")
    d.add_paragraph(
        "This letter is not a denial of coverage. We reserve the right to disclaim or limit coverage should investigation "
        "reveal facts that place this claim outside policy coverage. Please contact the assigned adjuster, "
        f"{c['adjuster']['name']}, at {c['adjuster']['phone']} with any questions.")
    d.add_paragraph(f"\nSincerely,\n{c['adjuster']['name']}\nClaims Department, {c['carrier']}")
    path = out / f"{cid}_coverage_position_letter.docx"
    d.save(path)
    return path


def release_docx(cid: str, c: dict, out: Path) -> Path:
    cl = c["claimant"]
    d = _doc(c["carrier"], "RELEASE OF ALL CLAIMS")
    d.add_paragraph(
        f"In consideration of the payment of ${c['settlement']['amount']:,.2f}, the receipt and sufficiency of which is "
        f"hereby acknowledged, {cl['name']} (“Releasor”) does hereby release and forever discharge "
        f"{c['insured']['name']} and {c['carrier']} (“Releasees”) from any and all claims, demands, and causes of "
        f"action arising out of the motor-vehicle collision that occurred on {c['loss']['date']} at {c['loss']['location']}.")
    d.add_paragraph(
        "Releasor acknowledges that this release covers all known and unknown injuries and represents a full and final "
        "settlement. This release is governed by the laws of the State of Texas.")
    d.add_paragraph("\n\n_______________________________          Date: ____________")
    d.add_paragraph(f"{cl['name']}, Releasor")
    path = out / f"{cid}_release_of_all_claims.docx"
    d.save(path)
    return path


def evaluation_docx(cid: str, c: dict, out: Path) -> Path:
    med = c["medical"]
    d = _doc(c["carrier"], "CLAIM EVALUATION & SETTLEMENT AUTHORITY REQUEST")
    d.add_paragraph(f"Claim: {cid}    Adjuster: {c['adjuster']['name']}    Reserves: "
                    f"${c['reserves']['bi']:,.0f} BI / ${c['reserves']['pd']:,.0f} PD")
    d.add_paragraph("Liability Analysis:")
    d.add_paragraph(c["loss"]["liability"] + " Insured presumed majority at fault; comparative negligence of 15–25% "
                    "attributable to claimant is arguable based on speed and following distance.")
    d.add_paragraph("Damages:")
    d.add_paragraph(f"Medical specials: ${med['billed']:,.2f}. Diagnoses: {'; '.join(med['diagnoses'])}. "
                    f"Plaintiff demand: ${c['demand']['amount']:,.2f}.")
    d.add_paragraph("Recommendation:")
    d.add_paragraph(f"Recommend settlement authority of ${c['settlement']['amount']:,.2f}, reflecting comparative fault and "
                    "litigation risk. Trial exposure estimated materially higher if liability is not apportioned.")
    path = out / f"{cid}_claim_evaluation.docx"
    d.save(path)
    return path


def estimate_xlsx(cid: str, c: dict, out: Path) -> Path:
    est = c["estimate"]
    v = c["vehicle"]
    wb = Workbook()
    ws = wb.active
    ws.title = "Estimate"
    hdr = Font(bold=True, color="FFFFFF")
    fill = PatternFill("solid", fgColor="142850")
    ws["A1"] = f"Repair Estimate — {cid}"
    ws["A1"].font = Font(bold=True, size=13)
    ws["A2"] = f"Vehicle: {v['year']} {v['make']} {v['model']} (VIN {v['vin']})"
    ws["A3"] = f"Insurer: {c['carrier']}"
    for col, h in zip("AB", ["Line Item", "Amount (USD)"]):
        cell = ws[f"{col}5"]
        cell.value = h
        cell.font = hdr
        cell.fill = fill
    rows = [("Parts", est["parts"]), ("Labor", est["labor"]), ("Paint & Materials", est["paint"]),
            ("Tax", est["tax"]), ("Subtotal (Total)", est["total"]), ("Less Deductible", -est["deductible"]),
            ("Net Payable", est["net"])]
    for i, (k, val) in enumerate(rows, start=6):
        ws[f"A{i}"] = k
        ws[f"B{i}"] = val
        ws[f"B{i}"].number_format = "$#,##0.00"
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 16
    path = out / f"{cid}_repair_estimate.xlsx"
    wb.save(path)
    return path
